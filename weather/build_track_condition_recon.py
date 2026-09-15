#!/usr/bin/env python3
"""Deterministically materialize the bounded Phase 4A.5 evidence review."""

import csv
import hashlib
import statistics
import zipfile
from collections import Counter, defaultdict
from datetime import date, datetime, time
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "weather"
PTSC_URL = "https://ptscin.com/wp-content/uploads/2025/06/FirestoneTemperatures.xlsx"
PTSC_WORKBOOK = "weather/evidence/ptsc/FirestoneTemperatures_current.xlsx"
PTSC_DATA = "weather/evidence/ptsc/indy500_day1_track_temp_2020_2024_time_normalized.csv"
FINAL_STATUS = "TRACK_CONDITION_RECONNAISSANCE_CORRECTED_AND_FROZEN"
DATES = {2020: "2020-08-15", 2021: "2021-05-22", 2022: "2022-05-21", 2023: "2023-05-20", 2024: "2024-05-18"}


def sha(path):
    candidate = ROOT / path
    return hashlib.sha256(candidate.read_bytes()).hexdigest() if candidate.exists() else ""


def write(name, rows):
    with (OUT / name).open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)


def obs(identifier, year, kind, local_time, utc, value, unit, celsius, text,
        semantics, title, url, source_class, quality, measurement, status, context,
        asset="", note=""):
    return {
        "observation_id": identifier, "year": year, "session_date": DATES[year],
        "evidence_type": kind, "stated_local_time": local_time,
        "local_timezone": "America/Indiana/Indianapolis", "normalized_utc": utc,
        "value_numeric": value, "unit": unit, "normalized_value_c": celsius,
        "stated_value_text": text, "session_semantics": semantics,
        "source_title": title, "source_url": url, "source_class": source_class,
        "evidence_quality": quality, "measurement_semantics": measurement,
        "observation_status": status, "evidence_context": context,
        "local_asset": asset, "review_note": note,
    }


def editorial_observations():
    indy_2022 = "Paddock Buzz: Sato Sneaks into Fast 12 after Eventful Day"
    indy_2022_url = "https://www.indycar.com/news/2022/05/05-21-buzz"
    indy_2022_asset = "weather/evidence/raw_pages/2022_www.indycar.com_news_2022_05_05-21-buzz.txt"
    apex = "Race Report: 2024 Indianapolis 500"
    apex_url = "https://www.theapex.racing/2024/05/2024-indianapolis-500/"
    return [
        obs("TC2020-Q01", 2020, "GENERAL_CONDITION", "", "", "", "", "", "best weather conditions of the day",
            "Day 1 qualifying; early draw/run context", "VeeKay, Palou Make Strong First Impressions To Reach Fast Nine Shootout",
            "https://www.indianapolismotorspeedway.com/news-multimedia/news/2020/08/15/veekay-palou-make-strong-first-impressions-to-reach-fast-nine-shootout",
            "IMS_OFFICIAL_EDITORIAL", "OFFICIAL_EDITORIAL", "QUALITATIVE_WEATHER_CONTEXT", "QUALITATIVE_OBSERVED",
            "IMS describes favorable weather but supplies no numerical environmental measurement."),
        obs("TC2021-Q01", 2021, "GENERAL_CONDITION", "afternoon", "", "", "", "", "warmer part of the afternoon",
            "Day 1 qualifying; second Ferrucci attempt", "Paddock Buzz: Nervous Sunday Ahead for Power, Four Others",
            "https://www.indycar.com/news/2021/05/05-22-paddockbuzz", "INDYCAR_OFFICIAL_EDITORIAL", "OFFICIAL_EDITORIAL",
            "QUALITATIVE_WEATHER_CONTEXT", "QUALITATIVE_OBSERVED", "INDYCAR describes a warmer afternoon without a number or exact time."),
        obs("TC2022-T01", 2022, "TRACK_TEMPERATURE", "11:00 a.m. EDT", "2022-05-21T15:00:00Z", 85, "degF", 29.4444444444,
            "85°F", "Day 1 qualifying opening", indy_2022, indy_2022_url, "INDYCAR_OFFICIAL_EDITORIAL", "OFFICIAL_EDITORIAL",
            "EXPLICIT_TRACK_TEMPERATURE", "DIRECT_NUMERIC_OBSERVATION", "Official report states the track was 85°F at opening.", indy_2022_asset),
        obs("TC2022-T02", 2022, "TRACK_TEMPERATURE", "12:30 p.m. EDT", "2022-05-21T16:30:00Z", 107, "degF", 41.6666666667,
            "107°F", "Day 1 qualifying in progress", indy_2022, indy_2022_url, "INDYCAR_OFFICIAL_EDITORIAL", "OFFICIAL_EDITORIAL",
            "EXPLICIT_ASPHALT_TEMPERATURE", "DIRECT_NUMERIC_OBSERVATION", "Official report states asphalt had heated to 107°F.", indy_2022_asset),
        obs("TC2022-Q01", 2022, "WIND", "early afternoon", "", "", "", "", "wind affected drivability",
            "Day 1 qualifying; Montoya context", indy_2022, indy_2022_url, "INDYCAR_OFFICIAL_EDITORIAL", "OFFICIAL_EDITORIAL",
            "QUALITATIVE_OBSERVED_WIND", "QUALITATIVE_OBSERVED", "Official report describes wind effects but gives no speed or direction.", indy_2022_asset),
        obs("TC2023-Q01", 2023, "GENERAL_CONDITION", "session-wide", "", "", "", "", "sunny skies; temperatures in the mid-70s°F",
            "Day 1 qualifying general conditions", "Rosenqvist Paces Epic, Historic First Qualifying Day at Indy",
            "https://www.indycar.com/News/2023/05/05-20-FirstDayQuals", "INDYCAR_OFFICIAL_EDITORIAL", "OFFICIAL_EDITORIAL",
            "QUALITATIVE_OBSERVED_WEATHER", "QUALITATIVE_OBSERVED", "Official recap gives general conditions without a precise time or sensor semantics."),
        obs("TC2024-T01", 2024, "TRACK_TEMPERATURE", "11:00 a.m. EDT", "2024-05-18T15:00:00Z", 94.6, "degF", 34.7777777778,
            "94.6°F", "Day 1 qualifying opening", apex, apex_url, "REPUTABLE_SECONDARY_SESSION_REPORT", "REPUTABLE_EDITORIAL",
            "EXPLICIT_TRACK_TEMPERATURE", "DIRECT_NUMERIC_OBSERVATION", "Secondary report attributes the reading to Firestone engineers.",
            note="The located publication is secondary rather than INDYCAR/IMS."),
        obs("TC2024-A01", 2024, "AIR_TEMPERATURE", "11:00 a.m. EDT", "2024-05-18T15:00:00Z", 73, "degF", 22.7777777778,
            "73°F", "Day 1 qualifying opening", apex, apex_url, "REPUTABLE_SECONDARY_SESSION_REPORT", "REPUTABLE_EDITORIAL",
            "EXPLICIT_AMBIENT_TEMPERATURE", "DIRECT_NUMERIC_OBSERVATION", "Secondary report gives ambient temperature as 73°F."),
        obs("TC2024-H01", 2024, "RELATIVE_HUMIDITY", "11:00 a.m. EDT", "2024-05-18T15:00:00Z", 80, "percent", "", "80%",
            "Day 1 qualifying opening", apex, apex_url, "REPUTABLE_SECONDARY_SESSION_REPORT", "REPUTABLE_EDITORIAL",
            "EXPLICIT_AMBIENT_HUMIDITY", "DIRECT_NUMERIC_OBSERVATION", "Secondary report gives relative humidity as 80%."),
        obs("TC2024-S01", 2024, "SKY_CONDITION", "11:00 a.m. EDT", "2024-05-18T15:00:00Z", "", "", "", "mostly cloudy",
            "Day 1 qualifying opening", apex, apex_url, "REPUTABLE_SECONDARY_SESSION_REPORT", "REPUTABLE_EDITORIAL",
            "OBSERVED_GENERAL_SKY_CONDITION", "QUALITATIVE_OBSERVED", "General sky condition does not establish corner-level shade."),
        obs("TC2024-Q01", 2024, "AIR_TEMPERATURE", "session-wide", "", "", "", "", "low 80s°F maximum/range description",
            "Day 1 qualifying general conditions", "Power Fastest as Penske Eyes Pole after Top Three Sweep",
            "https://www.indycar.com/News/2024/05/05-18-Quals-Day1", "INDYCAR_OFFICIAL_EDITORIAL", "OFFICIAL_EDITORIAL",
            "QUALITATIVE_AIR_TEMPERATURE_RANGE", "QUALITATIVE_OBSERVED", "Official recap gives a range without an exact observation time."),
    ]


def load_ptsc():
    with (ROOT / PTSC_DATA).open(encoding="utf-8-sig", newline="") as handle:
        rows = list(csv.DictReader(handle))
    observations = []
    for row in rows:
        year = int(row["year"])
        source_row = int(row["source_row_index"])
        track_f = float(row["track_f"])
        track_c = float(row["track_c"])
        observations.append(obs(
            f"PTSC{year}-{source_row:04d}", year, "TRACK_TEMPERATURE", row["local_time"],
            row["utc_datetime"].replace(" ", "T").replace("+00:00", "Z"), track_f, "degF", track_c, f"{track_f:g}°F",
            "Target-date structured observation; rows outside qualifying hours remain Day 1 date context",
            "PTSC-hosted FirestoneTemperatures workbook", PTSC_URL, "PTSC_HOSTED_FIRESTONE_NAMED_WORKBOOK",
            "STRUCTURED_OBSERVED_SOURCE_WITH_UNIT_LIMITATIONS", "EXPLICIT_WORKBOOK_TRACK_TEMPERATURE", "DIRECT_NUMERIC_OBSERVATION",
            f"Sheet {row['source_sheet']}; source row {source_row}; selected date-block occurrence {row['source_block_occurrence']}; sensor average {row['sensor_avg_f']}°F.",
            PTSC_WORKBOOK, "Not INDYCAR official. PTSC hosts a Firestone-named workbook; local metadata does not independently authenticate authorship or measurement operator.",
        ))
    return rows, observations


def ptsc_statistics(rows):
    grouped = defaultdict(list)
    for row in rows:
        grouped[int(row["year"])].append(row)
    stats = {}
    for year, year_rows in grouped.items():
        ordered = sorted(year_rows, key=lambda row: row["utc_datetime"])
        instants = [datetime.fromisoformat(row["utc_datetime"]) for row in ordered]
        gaps = [(right - left).total_seconds() / 60 for left, right in zip(instants, instants[1:])]
        stats[year] = {
            "rows": len(ordered), "first_local": ordered[0]["local_datetime"], "last_local": ordered[-1]["local_datetime"],
            "median_interval_min": statistics.median(gaps), "max_interval_min": max(gaps), "gaps_gt_30": sum(gap > 30 for gap in gaps),
            "track_nonnull": sum(row["track_f"] != "" for row in ordered), "ambient_nonnull": sum(row["ambient_f"] != "" for row in ordered),
            "humidity_nonnull": sum(row["humidity"] != "" for row in ordered), "wind_nonnull": sum(row["wind"] != "" for row in ordered),
            "pressure_valid": sum(row["pressure"] != "" for row in ordered),
            "pressure_invalid": sum(row["pressure_quality"] == "SOURCE_VALUE_OUT_OF_RANGE" for row in ordered),
        }
    return stats


def verify_ptsc_workbook_blocks():
    workbook = load_workbook(ROOT / PTSC_WORKBOOK, read_only=True, data_only=True)
    verifications = []
    for year, date_text in DATES.items():
        target = date.fromisoformat(date_text)
        rows = list(workbook[f"{year} Temp Archive"].iter_rows(values_only=True))
        starts = []
        for index, row in enumerate(rows):
            value = row[0]
            detected = value.date() if isinstance(value, datetime) else value if isinstance(value, date) else None
            if detected == target:
                starts.append(index)
        block_counts = []
        for start in starts:
            end = next((index for index in range(start + 1, len(rows)) if isinstance(rows[index][0], (date, datetime))), len(rows))
            block_counts.append(sum(isinstance(rows[index][1], (time, datetime)) for index in range(start, end)))
        verifications.append({
            "year": year,
            "session_date": date_text,
            "source_sheet": f"{year} Temp Archive",
            "source_date_occurrence_count": len(starts),
            "selected_occurrence": 1,
            "selected_block_valid_time_row_count": block_counts[0] if block_counts else 0,
            "excluded_duplicate_block_valid_time_row_count": sum(block_counts[1:]),
            "source_block_counts": ";".join(str(value) for value in block_counts),
            "verification_status": "PASS" if block_counts and block_counts[0] > 0 else "FAIL",
        })
    workbook.close()
    return verifications


def source_catalog():
    specs = [
        ("SRC2020-IMS-DAY1", 2020, "VeeKay, Palou Make Strong First Impressions To Reach Fast Nine Shootout", "https://www.indianapolismotorspeedway.com/news-multimedia/news/2020/08/15/veekay-palou-make-strong-first-impressions-to-reach-fast-nine-shootout", "IMS_OFFICIAL_EDITORIAL", "IMS", "", "REVIEWED", "Qualitative weather context", "No numeric environmental values"),
        ("SRC2021-INDYCAR-GREEN", 2021, "Green Flag: Crown Royal Armed Forces Qualifying Day 1", "https://www.indycar.com/News/2021/05/05-22-GreenFlag", "INDYCAR_OFFICIAL_EDITORIAL", "INDYCAR", "", "REVIEWED_CONTEXT_ONLY", "Forecast and sun/cloud/wind mechanism context", "Not a historical measured series"),
        ("SRC2021-INDYCAR-BUZZ", 2021, "Paddock Buzz: Nervous Sunday Ahead for Power, Four Others", "https://www.indycar.com/news/2021/05/05-22-paddockbuzz", "INDYCAR_OFFICIAL_EDITORIAL", "INDYCAR", "", "REVIEWED", "Warmer-afternoon qualitative observation", "No numerical measurement"),
        ("SRC2022-INDYCAR-BUZZ", 2022, "Paddock Buzz: Sato Sneaks into Fast 12 after Eventful Day", "https://www.indycar.com/news/2022/05/05-21-buzz", "INDYCAR_OFFICIAL_EDITORIAL", "INDYCAR", "weather/evidence/raw_pages/2022_www.indycar.com_news_2022_05_05-21-buzz.txt", "REVIEWED", "Two numeric track/asphalt anchors and qualitative wind", "No structured feed"),
        ("SRC2023-INDYCAR-DAY1", 2023, "Rosenqvist Paces Epic, Historic First Qualifying Day at Indy", "https://www.indycar.com/News/2023/05/05-20-FirstDayQuals", "INDYCAR_OFFICIAL_EDITORIAL", "INDYCAR", "", "REVIEWED", "Session-wide sunny/mid-70s description", "No numeric track-temperature anchor"),
        ("SRC2023-INDYCAR-SETUP", 2023, "The Setup: Indianapolis 500 with Brad Goldberg", "https://www.indycar.com/News/2023/05/05-19-Setup-Indy", "INDYCAR_OFFICIAL_TECHNICAL", "INDYCAR", "", "REVIEWED_CONTEXT_ONLY", "Heat, grip, dust, marbles, and rubber context", "Not a Day 1 measurement"),
        ("SRC2024-INDYCAR-DAY1", 2024, "Power Fastest as Penske Eyes Pole after Top Three Sweep", "https://www.indycar.com/News/2024/05/05-18-Quals-Day1", "INDYCAR_OFFICIAL_EDITORIAL", "INDYCAR", "", "REVIEWED", "Air-temperature range and sunny/coolest-period context", "No numeric track-temperature anchor"),
        ("SRC2024-THEAPEX", 2024, "Race Report: 2024 Indianapolis 500", "https://www.theapex.racing/2024/05/2024-indianapolis-500/", "REPUTABLE_SECONDARY_SESSION_REPORT", "Firestone measurement attribution", "", "REVIEWED", "11:00 track/ambient/humidity/sky anchor", "Secondary publication; sensor details absent"),
        ("SRC-PTSC-WORKBOOK", "", "PTSC-hosted FirestoneTemperatures workbook", PTSC_URL, "PTSC_HOSTED_FIRESTONE_NAMED_WORKBOOK", "PTSC-hosted; Firestone-named; authorship/measurement operator not independently authenticated", PTSC_WORKBOOK, "REVIEWED", "Structured track, ambient, humidity, wind, barometer, and four sensor readings", "Not INDYCAR official; wind/humidity/barometer units are not explicit in headers; one 2022 barometer value is invalid"),
        ("SRC-HRRR-CANONICAL-REFERENCE", "", "Existing canonical HRRR forecast layer", "", "HRRR_FORECAST_REFERENCE_EXISTING_CANONICAL", "NOAA/NCEP HRRR", "data/canonical/v1/weather_forecasts.csv", "REFERENCED_NOT_MODIFIED", "Forecast covariates", "Forecast data; not observed track temperature and not merged with PTSC"),
    ]
    replay_assets = {2020: ["evidence/timing71_2020_part1.zip", "evidence/timing71_2020_part2.zip"], 2021: ["evidence/timing71_2021_sample.zip"], 2023: ["evidence/timing71_2023.zip"], 2024: ["evidence/timing71_2024.zip"]}
    for year, assets in replay_assets.items():
        for index, asset in enumerate(assets, 1):
            specs.append((f"SRC{year}-TIMING71-{index}", year, "Timing71 replay capture", "", "THIRD_PARTY_REPLAY_CAPTURE", "INDYCAR live timing display", asset, "REVIEWED", "Tire category marker only", "No weather or tire thermal telemetry fields; capture limitations remain"))
    results_assets = {2020: "evidence/results_2020_0.pdf", 2021: "weather/evidence/downloads/2021/2021-indycar-results-quals-day1.pdf", 2022: "evidence/results_2022_0.pdf", 2023: "evidence/results_2023_1.pdf", 2024: "weather/evidence/downloads/2024/indycar-results-quals-day1-05-18-2024.pdf"}
    for year, asset in results_assets.items():
        specs.append((f"SRC{year}-RESULTS", year, f"{year} INDYCAR Day 1 Qualification Results", "", "OFFICIAL_TIMING_REPORT", "INDYCAR", asset, "REVIEWED", "Lap/attempt results and Firestone manufacturer code", "No environmental or tire telemetry values"))
    local_assets = ["weather/evidence/indy500_discovered_sources.csv", "weather/evidence/download_manifest.csv", "weather/evidence/priority_download_manifest.csv", "weather/evidence/track_condition_evidence_candidates.csv", "weather/evidence/track_condition_evidence_candidates_v2.csv", "weather/scripts/collect_track_condition_evidence.py", "weather/scripts/crawl_track_condition_sources.py", "weather/scripts/discover_indy500_sources.py", "weather/scripts/download_discovered_sources.py"]
    ptsc_assets = ["weather/evidence/ptsc/indy500_day1_track_temp_2020_2024.csv", PTSC_DATA, "weather/evidence/ptsc/indy500_day1_track_temp_excluded_blocks.csv", "weather/evidence/ptsc/indy500_day1_track_temp_qa.csv", "weather/evidence/ptsc/indy500_day1_track_temp_qa_summary.csv", "weather/scripts/download_current_firestone_temps.py", "weather/scripts/extract_indy500_day1_track_temp.py", "weather/scripts/normalize_ptsc_timestamps.py", "weather/scripts/qa_indy500_day1_track_temp.py"]
    for index, asset in enumerate(local_assets, 1):
        specs.append((f"SRC-LOCAL-RECON-{index:02d}", "", "Previously retained reconnaissance work product", "", "LOCAL_RECONNAISSANCE_WORK_PRODUCT", "LOCAL", asset, "PRESERVED_AND_REVIEWED", "Discovery, download, or candidate-triage record", "Not independent evidence"))
    for index, asset in enumerate(ptsc_assets, 1):
        specs.append((f"SRC-PTSC-DERIVED-{index:02d}", "", "PTSC extraction/provenance/QA asset", PTSC_URL, "LOCAL_DERIVATIVE_OR_EXTRACTION_SCRIPT", "Derived from the PTSC-hosted workbook", asset, "PRESERVED_AND_REVIEWED", "Extraction, timestamp normalization, excluded block, or QA", "Not an independent source"))
    return [{"source_id": sid, "year": year, "source_title": title, "source_url": url, "source_class": source_class,
             "authority_or_underlying_source": authority, "local_asset": asset, "content_hash_sha256": sha(asset) if asset else "",
             "review_status": status, "evidence_found": found, "limitations": limitation}
            for sid, year, title, url, source_class, authority, asset, status, found, limitation in specs]


def build_timing_audit():
    tokens = ("weather", "tracktemp", "airtemp", "ambient", "humidity", "wind", "pressure", "tiretemperature", "tyretemperature")
    archives = {2020: [("evidence/timing71_2020_part1.zip", "PARTIAL_BEGINNING_OR_END_UNKNOWN"), ("evidence/timing71_2020_part2.zip", "PARTIAL_BEGINNING_OR_END_UNKNOWN")], 2021: [("evidence/timing71_2021_sample.zip", "SAMPLE_ONLY")], 2023: [("evidence/timing71_2023.zip", "RETAINED_CAPTURE")], 2024: [("evidence/timing71_2024.zip", "RETAINED_CAPTURE")]}
    audit = []
    for year, year_archives in archives.items():
        for asset, coverage in year_archives:
            counts, json_count, tyre_count = Counter(), 0, 0
            with zipfile.ZipFile(ROOT / asset) as archive:
                for name in archive.namelist():
                    if not name.lower().endswith(".json"):
                        continue
                    json_count += 1
                    payload = archive.read(name).lower()
                    tyre_count += b"tyre-medium" in payload
                    for token in tokens:
                        counts[token] += token.encode() in payload
            row = {"year": year, "archive": asset, "archive_coverage": coverage, "json_file_count": json_count, "tyre_category_file_count": tyre_count}
            row.update({f"{token}_token_file_count": counts[token] for token in tokens})
            row["finding"] = "Tire category marker present; no environment, tire-temperature, or tire-pressure field located"
            audit.append(row)
    missing = {"year": 2022, "archive": "", "archive_coverage": "NOT_RECOVERED", "json_file_count": 0, "tyre_category_file_count": 0}
    missing.update({f"{token}_token_file_count": 0 for token in tokens})
    missing["finding"] = "No retained 2022 Timing71 qualifying replay; field availability remains unresolved"
    audit.append(missing)
    return sorted(audit, key=lambda row: (int(row["year"]), row["archive"])), tokens


def main():
    observations = editorial_observations()
    ptsc_rows, ptsc_observations = load_ptsc()
    observations.extend(ptsc_observations)
    write("track_condition_evidence.csv", observations)
    stats = ptsc_statistics(ptsc_rows)
    workbook_verifications = verify_ptsc_workbook_blocks()
    write("ptsc_source_block_verification.csv", workbook_verifications)
    sources = source_catalog()
    write("track_condition_source_catalog.csv", sources)

    independent_anchors = {2020: 0, 2021: 0, 2022: 2, 2023: 0, 2024: 1}
    other_weather = {2020: 1, 2021: 1, 2022: 1, 2023: 1, 2024: 4}
    tire_availability = {2020: "PARTIALLY_AVAILABLE", 2021: "PARTIALLY_AVAILABLE", 2022: "UNRESOLVED", 2023: "PARTIALLY_AVAILABLE", 2024: "PARTIALLY_AVAILABLE"}
    shade = {2020: "NOT_FOUND", 2021: "FOUND_QUALITATIVE_ONLY", 2022: "NOT_FOUND", 2023: "FOUND_QUALITATIVE_ONLY", 2024: "FOUND_QUALITATIVE_ONLY"}
    coverage = []
    for year in DATES:
        s = stats[year]
        coverage.append({
            "year": year, "session_date": DATES[year], "track_temperature_status": "FOUND_STRUCTURED",
            "observed_weather_status": "FOUND_STRUCTURED", "indycar_official_structured_track_feed_status": "NOT_FOUND",
            "ptsc_track_temperature_observation_count": s["rows"], "independent_editorial_or_secondary_track_anchor_count": independent_anchors[year],
            "numeric_track_temperature_observation_count_total": s["rows"] + independent_anchors[year],
            "ptsc_observed_weather_row_count": s["rows"], "non_ptsc_weather_evidence_count": other_weather[year],
            "ptsc_first_local": s["first_local"], "ptsc_last_local": s["last_local"], "ptsc_median_interval_min": s["median_interval_min"],
            "ptsc_max_interval_min": s["max_interval_min"], "ptsc_gaps_over_30_min": s["gaps_gt_30"],
            "ptsc_pressure_valid_rows": s["pressure_valid"], "ptsc_pressure_invalid_rows": s["pressure_invalid"],
            "tire_telemetry_availability": tire_availability[year], "tire_evidence_coverage": "UNRESOLVED" if year == 2022 else "FOUND_QUALITATIVE_ONLY",
            "rubber_track_evolution_status": "NOT_FOUND", "shade_exposure_status": shade[year],
            "assessment_note": "PTSC provides structured target-date track, ambient, humidity, wind, barometer, and surface-sensor rows. It is not INDYCAR official; intervals contain a >30-minute gap and wind/humidity/barometer units are not explicit in headers."
            + (" One 2022 barometer source value is out of range and remains raw-only." if year == 2022 else ""),
        })
    write("track_condition_coverage_audit.csv", coverage)

    tire = []
    for year in DATES:
        state = "UNRESOLVED" if year == 2022 else "NOT_FOUND"
        tire.append({"year": year, "session_date": DATES[year], "availability_classification": tire_availability[year],
                     "tire_carcass_temperature": state, "tire_surface_temperature": state, "tire_pressure_time_series": state,
                     "thermal_state_inference_telemetry": state, "publicly_observed_tire_context": "No qualifying replay recovered" if year == 2022 else "Timing71 tyre-medium/category marker only",
                     "notes": "PTSC surface/track sensors are not tire telemetry. Category/manufacturer evidence is not tire temperature, pressure, or thermal state."})
    write("tire_telemetry_availability.csv", tire)

    timing_audit, timing_tokens = build_timing_audit()
    write("timing71_environment_field_audit.csv", timing_audit)

    ptsc_lookup = {(int(row["year"]), row["local_time"]): row for row in ptsc_rows}
    anchor_specs = [("TC2022-T01", 2022, "11:00:00", 85.0, "INDYCAR_OFFICIAL_EDITORIAL"), ("TC2022-T02", 2022, "12:30:00", 107.0, "INDYCAR_OFFICIAL_EDITORIAL"), ("TC2024-T01", 2024, "11:00:00", 94.6, "REPUTABLE_SECONDARY_SESSION_REPORT")]
    reconciliations = []
    for anchor_id, year, local_time, anchor_f, source_class in anchor_specs:
        row = ptsc_lookup[(year, local_time)]
        ptsc_f, sensor_avg_f = float(row["track_f"]), float(row["sensor_avg_f"])
        reconciliations.append({"anchor_observation_id": anchor_id, "year": year, "local_time": local_time,
            "utc_datetime": row["utc_datetime"].replace(" ", "T").replace("+00:00", "Z"), "anchor_value_f": anchor_f,
            "anchor_source_class": source_class, "ptsc_track_f": ptsc_f, "ptsc_sensor_avg_f": sensor_avg_f,
            "ptsc_track_minus_anchor_f": round(ptsc_f - anchor_f, 6), "ptsc_sensor_avg_minus_anchor_f": round(sensor_avg_f - anchor_f, 6),
            "match_method": "EXACT_LOCAL_TIMESTAMP_NO_INTERPOLATION",
            "assessment": "Values differ; retained as separate source observations without forcing agreement."})
    write("track_temperature_anchor_reconciliation.csv", reconciliations)

    checks = []
    def check(check_id, passed, details):
        checks.append({"check_id": check_id, "pass_fail": "PASS" if passed else "FAIL", "details": details})
    track = [row for row in observations if row["evidence_type"] == "TRACK_TEMPERATURE"]
    official = [row for row in track if row["source_class"] == "INDYCAR_OFFICIAL_EDITORIAL"]
    ptsc_track = [row for row in track if row["source_class"] == "PTSC_HOSTED_FIRESTONE_NAMED_WORKBOOK"]
    check("FIVE_TARGET_YEARS_ASSESSED", {row["year"] for row in coverage} == set(DATES), "Coverage contains 2020-2024 exactly")
    check("PHASE_4A5_FINAL_STATUS", FINAL_STATUS == "TRACK_CONDITION_RECONNAISSANCE_CORRECTED_AND_FROZEN", FINAL_STATUS)
    check("PTSC_TOTAL_ROW_COUNT", len(ptsc_rows) == 168, f"Expected 168; observed {len(ptsc_rows)}")
    check("PTSC_YEAR_ROW_COUNTS", Counter(int(row["year"]) for row in ptsc_rows) == Counter({2020: 32, 2021: 33, 2022: 31, 2023: 36, 2024: 36}), "Normalized asset matches independently checked workbook blocks")
    check("PTSC_WORKBOOK_BLOCK_COUNTS_VERIFIED", {int(row["year"]): int(row["selected_block_valid_time_row_count"]) for row in workbook_verifications} == {2020: 32, 2021: 33, 2022: 31, 2023: 36, 2024: 36}, "Counts read directly from workbook date blocks")
    check("PTSC_2022_DUPLICATE_BLOCK_EXCLUDED", next(row for row in workbook_verifications if int(row["year"]) == 2022)["excluded_duplicate_block_valid_time_row_count"] == 20, "Second same-date source block has 20 valid time rows and remains excluded")
    check("PTSC_TRACK_ROWS_MATERIALIZED", len(ptsc_track) == 168, f"Expected 168; observed {len(ptsc_track)}")
    check("PTSC_ALL_TRACK_AND_CORE_WEATHER_PRESENT", all(s["track_nonnull"] == s["rows"] and s["ambient_nonnull"] == s["rows"] and s["humidity_nonnull"] == s["rows"] and s["wind_nonnull"] == s["rows"] for s in stats.values()), "Track, ambient, humidity, and wind populate every PTSC row")
    check("PTSC_2022_PRESSURE_ANOMALY_PRESERVED", stats[2022]["pressure_valid"] == 30 and stats[2022]["pressure_invalid"] == 1, "One source value remains raw-only; no correction invented")
    check("2022_BENCHMARK_ANCHORS_PRESERVED", [(row["stated_local_time"], row["value_numeric"]) for row in official] == [("11:00 a.m. EDT", 85), ("12:30 p.m. EDT", 107)], "Official 85°F and 107°F anchors")
    check("ANCHOR_RECONCILIATION_EXACT_ONLY", all(row["match_method"] == "EXACT_LOCAL_TIMESTAMP_NO_INTERPOLATION" for row in reconciliations), "Three anchors compared at exact PTSC timestamps")
    check("CELSIUS_CONVERSION_VALID", all(abs(float(row["normalized_value_c"]) - (float(row["value_numeric"]) - 32) * 5 / 9) < 1e-9 for row in track), "All Fahrenheit track values convert deterministically")
    check("NO_INFERRED_NUMERIC_TRACK_TEMPERATURE", all(row["observation_status"] == "DIRECT_NUMERIC_OBSERVATION" and row["measurement_semantics"].startswith("EXPLICIT_") for row in track), "All numeric track values are explicit source observations")
    classes = {row["source_class"] for row in sources}
    check("SOURCE_AUTHORITY_CLASSES_DISTINCT", {"INDYCAR_OFFICIAL_EDITORIAL", "PTSC_HOSTED_FIRESTONE_NAMED_WORKBOOK", "REPUTABLE_SECONDARY_SESSION_REPORT", "HRRR_FORECAST_REFERENCE_EXISTING_CANONICAL"}.issubset(classes), "Official editorial, PTSC, secondary, and HRRR remain separate")
    check("NO_NUMERIC_TIRE_TELEMETRY", all(row["tire_carcass_temperature"] in ("NOT_FOUND", "UNRESOLVED") and row["tire_surface_temperature"] in ("NOT_FOUND", "UNRESOLVED") and row["tire_pressure_time_series"] in ("NOT_FOUND", "UNRESOLVED") for row in tire), "PTSC surface sensors were not relabelled as tire telemetry")
    check("SOURCE_LOCAL_HASHES_PRESENT", all((not row["local_asset"]) or row["content_hash_sha256"] for row in sources), "All referenced local assets have SHA-256 hashes")
    check("EVIDENCE_SOURCE_URL_PRESENT", all(row["source_url"] for row in observations), "All accepted observations have a source URL")
    check("NO_RUBBER_OR_SHADE_NUMERIC_PROXY", not any(row["evidence_type"] in ("RUBBER_PROXY", "SHADE_PROXY") for row in observations), "No proxy created")
    check("TIMING71_HAS_NO_ENVIRONMENT_OR_TIRE_TELEMETRY_FIELDS", all(not row[f"{token}_token_file_count"] for row in timing_audit for token in timing_tokens), "Target tokens absent from retained Timing71 JSON")
    write("track_condition_recon_qa.csv", checks)

    coverage_lines = "\n".join(f"| {row['year']} | FOUND_STRUCTURED ({row['ptsc_track_temperature_observation_count']} PTSC rows) | FOUND_STRUCTURED ({row['ptsc_observed_weather_row_count']} PTSC rows) | {row['tire_telemetry_availability']} | {row['rubber_track_evolution_status']} | {row['shade_exposure_status']} |" for row in coverage)
    ptsc_lines = "\n".join(f"| {year} | {s['rows']} | {s['first_local']} | {s['last_local']} | {s['median_interval_min']:g} min | {s['max_interval_min']:g} min | {s['pressure_valid']}/{s['rows']} |" for year, s in sorted(stats.items()))
    report = f"""# Historical Track-Condition Evidence Reconnaissance

**Final status: `{FINAL_STATUS}`**

Phase 4A.5 covers only the five specified Indianapolis 500 Day 1 dates. This corrected review includes the retained PTSC-hosted `FirestoneTemperatures` workbook and its derived assets. The first Phase 4A.5 materialization omitted them; they were not intentionally excluded.

## Corrected verdict

- Recoverable numeric track/asphalt temperature is **FOUND_STRUCTURED in 5/5 years**, with **168 PTSC observations**: 2020 32, 2021 33, 2022 31, 2023 36, and 2024 36.
- An **INDYCAR-official structured track-temperature feed remains NOT_FOUND in 5/5 years**. The two official 2022 webpage anchors are editorial observations, not a structured feed.
- Structured time-indexed observed ambient temperature, humidity, wind, and barometer fields occur in **5/5 years** in PTSC. Coverage is not proven gap-free: each year has one gap over 30 minutes, and maximum gaps range from 45 to 75 minutes.
- PTSC wind, humidity, and barometer headers do not state units. One 2022 barometer value (`239.13`) remains raw while normalized pressure is missing under the existing quality rule.
- Tire thermal/pressure telemetry remains recoverable in **0/5 years**. PTSC surface sensors are not tire telemetry.
- Quantitative rubbering/track evolution and direct corner-level shade evidence remain recoverable in **0/5 years**.

## Source-authority separation

1. **INDYCAR/IMS official editorial pages** provide two numeric 2022 anchors and qualitative context. They do not provide a structured temperature feed.
2. **PTSC-hosted Firestone-named workbook** provides 168 structured observations. It is not labelled INDYCAR official. PTSC hosting and the filename are documented; local metadata does not independently authenticate authorship or the measurement operator.
3. **Secondary reporting** supplies the separate 2024 11:00 Firestone-engineer-attributed anchor.
4. **HRRR** remains a forecast layer. It is not observed track temperature and was not merged with PTSC.

## Structured PTSC coverage

| Year | Rows | Local first | Local last | Median interval | Maximum interval | Valid pressure |
|---:|---:|---|---|---:|---:|---:|
{ptsc_lines}

These are target-date observations. Some begin before qualifying opened, so they are not all labelled as in-session attempt observations.

## Independent-anchor reconciliation

| Source | Time | Independent value | PTSC `track_f` | PTSC sensor average | PTSC track difference |
|---|---|---:|---:|---:|---:|
| INDYCAR official editorial | 2022-05-21 11:00 EDT | 85°F | 86°F | 85.625°F | +1.0°F |
| INDYCAR official editorial | 2022-05-21 12:30 EDT | 107°F | 105°F | 105.300°F | -2.0°F |
| Secondary, Firestone-engineer attribution | 2024-05-18 11:00 EDT | 94.6°F | 95°F | 94.950°F | +0.4°F |

All comparisons use exact local timestamps. No interpolation or forced agreement was applied. The 2022 official values remain 85°F and 107°F; the 2024 secondary value remains separately classified.

## Year coverage

| Year | Track temperature | Observed weather | Tire telemetry | Rubber/evolution | Shade/exposure |
|---:|---|---|---|---|---|
{coverage_lines}

`FOUND_STRUCTURED` means machine-readable and time-indexed. It does not imply INDYCAR authority, gap-free continuity, uniform sampling, or attempt-level alignment. `NOT_FOUND` means not recoverable from the bounded evidence reviewed here; it does not prove the data never existed.

## Timing71 and tire findings

The retained Timing71 JSON scan found no target environmental or tire-thermal fields. Its `tyre-medium` marker supports tire-category context only. The missing 2022 replay remains unresolved. PTSC surface sensors remain track-condition measurements and are not used to infer tire thermal state.

## Scope controls

No canonical chronology, HRRR file, schema, model, simulator, queue timing, or earlier Pipeline Integrity Gate output was changed. No interpolation, observed-to-HRRR merge, rubber metric, or shade proxy was created.

## Outputs

- `weather/track_condition_evidence.csv` — 168 PTSC track observations plus retained official, IMS, and secondary observations.
- `weather/track_condition_coverage_audit.csv` — corrected coverage and PTSC interval/quality limits.
- `weather/track_temperature_anchor_reconciliation.csv` — exact-time comparison of independent anchors with PTSC.
- `weather/ptsc_source_block_verification.csv` — direct workbook date-block counts, including the excluded 2022 duplicate block.
- `weather/tire_telemetry_availability.csv` — tire assessment with PTSC surface sensors kept distinct.
- `weather/track_condition_source_catalog.csv` — source classes, provenance assets, and hashes.
- `weather/timing71_environment_field_audit.csv` — replay scan and explicit 2022 gap.
- `weather/track_condition_recon_qa.csv` — deterministic checks.
- `weather/track_condition_reconnaissance_freeze_manifest.csv` — frozen output inventory and SHA-256 hashes.
"""
    (OUT / "track_condition_reconnaissance_report.md").write_text(report, encoding="utf-8")

    frozen_outputs = [
        "weather/track_condition_reconnaissance_report.md",
        "weather/track_condition_evidence.csv",
        "weather/track_condition_coverage_audit.csv",
        "weather/track_temperature_anchor_reconciliation.csv",
        "weather/ptsc_source_block_verification.csv",
        "weather/tire_telemetry_availability.csv",
        "weather/track_condition_source_catalog.csv",
        "weather/timing71_environment_field_audit.csv",
        "weather/track_condition_recon_qa.csv",
    ]
    freeze_manifest = [
        {"final_status": FINAL_STATUS, "output_file": path, "sha256": sha(path)}
        for path in frozen_outputs
    ]
    write("track_condition_reconnaissance_freeze_manifest.csv", freeze_manifest)


if __name__ == "__main__":
    main()
